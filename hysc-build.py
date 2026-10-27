import re
import json
import fire
from glob import glob
from bs4 import BeautifulSoup
from docx import Document, oxml
from docx.oxml.ns import qn
from docx.shared import Pt
import docx.enum.text as doc_t
import difflib as df

punc_tail = '，。：；？“”‘’「」○'
re_cjk_b = re.compile(r'[\U00020000-\U0002a6d6]')
re_cjk_c = re.compile(r'[\U0002a700-\U0002fa1d\U000e0100-\U000e0176]')
re_punc = re.compile('[，。：；？“”‘’「」○《》〈〉．、▲]')
re_kp_num = re.compile(r'^(\d+)(\.\d+)')
re_p_tag = re.compile('^【[^】]+】')
re_p_tag2 = re.compile('^([疏鈔經钞经]序?)[：「]')
re_p_tags = re.compile('【[^】]+】')

re_jing = re.compile('^經序?：')
re_kuo_kuo = re.compile('([，、])([』」])([『「])')
re_kuo_embed = re.compile(r'「[^「」]+[『』\u3000，。；][^「」]*」')
re_kuo_dup = re.compile('「[^「」『』]{0,5}『[^「」『』]+』[等者]?」')
re_kuo_long = re.compile('「[^「」]{16,100}」')
re_kuo_sign = re.compile('[「」『』]')
re_kuo_end = re.compile('[，：；、]”$')
re_kp_level = re.compile(r'^\d+')
re_tag1 = re.compile('“([^“”「」，][^“”「」]+)”')
re_tag2 = re.compile('[言引▲]+([^者。”」，][^者。”」]+?)等?[，。]?[者下中]')
re_tag3 = re.compile('^[▲“]*([^▲言引者。“”」，][^者。“”」]+?)等?[者下中]')
re_term = re.compile('^言.+者$')

re_diff_eq = re.compile(r'(= .+\n)+')
re_diff_ad = re.compile(r'(\+ .+\n){2,9999}')
re_diff_rm = re.compile(r'(- .+\n){2,9999}')
re_cmp8 = re.compile(r'[\s\n=+-]')

data = dict(p_n=0, pu=[])


def diff_text(a, b):
    """比较两个文本，得到差异文本"""
    t = df.Differ().compare(a, b)
    t = ''.join(list(re.sub('^ ', '=', s) + '\n' for s in t))
    t = re_diff_eq.sub(lambda m: '  %s\n' % re_cmp8.sub('', m.group()), t)
    t = re_diff_ad.sub(lambda m: '+ %s\n' % re_cmp8.sub('', m.group()), t)
    t = re_diff_rm.sub(lambda m: '- %s\n' % re_cmp8.sub('', m.group()), t)
    return t


def read_kp(soup, kp_map):
    """
    读取科判树
    :param soup: HTML根节点
    :param kp_map: 存储科判数据
    """
    kp_root = soup.find('div', id='tree-menu-left-kp')
    kp_map['id'] = 0
    read_kp_ul(kp_root.find('ul', class_='jstree-children'), kp_map, kp_map['tree'], [], 0)


def read_kp_ul(ul, kp_map, siblings, base_nums, level):
    """
    读取科判 ul 节点
    :param ul: jstree-children 类的节点
    :param kp_map: 存储科判数据
    :param siblings: 存储直接子节点
    :param base_nums: 当前节点的各级科判序号
    :param level: 层次，根节点为0
    :return: 各级子节点的个数
    """
    if ul is None:
        return 0
    count = 0
    for li in ul.find_all('li', class_='jstree-node', recursive=False):
        a = li.find('a', class_='jstree-anchor', recursive=False)
        target, file = a.attrs['target_content'], a.attrs['html_file']
        k = re.search(r'\d+', file).group()
        id2node = kp_map['id2node'][k] = kp_map['id2node'].get(k, {})
        node_id = target[1:]
        assert target.startswith('#') and file[-5:] == '.html'
        assert node_id not in id2node
        kp_map['id'] += 1
        node = dict(target=target[1:], text=a.text.strip(), id=kp_map['id'], children=[])
        nums = base_nums
        if '.' in node['text']:
            # assert re.match(r'^\d+\.\d+ ', node['text'])
            # assert int(node['text'].split('.')[0]) == level + 1
            node['num'] = 1 + max([0] + [s.get('num', 0) for s in siblings])
            nums = node['nums'] = base_nums + [node['num']]
            num_txt = ('…' if len(nums) > 5 else '') + '.'.join([str(s) for s in nums[-5:]])
            node['text'] = re_kp_num.sub(lambda m: m.group(1) + ':' + num_txt, node['text'])
        node2 = id2node[node_id] = dict(node)
        [node2.pop(s, 0) for s in ['children', 'num', 'nums']]
        siblings.append(node)
        sub_n = read_kp_ul(li.find('ul', class_='jstree-children', recursive=False),
                           kp_map, node['children'], nums, level + 1)
        count += 1 + sub_n
        if not sub_n:
            node.pop('children')
    return count


def add_head_level(kp_map, book, juan):
    """
    对疏钞HTML中的科判段落节点添加编号文本
    :param kp_map: 科判数据
    :param book: 疏钞HTML的主内容soup对象
    :param juan: 疏钞的卷号文本
    """
    id2node = kp_map['id2node'][juan]
    for node in list(id2node.values()):
        a = book.select(f"a[name={node['target']}]")
        if a:
            head = a[0]
            while head.name != 'p' or 'head' not in head.attrs['class']:
                head = head.next_element
            assert head.text in node['text']
            if head.text != node['text']:
                head.string = node['text']
            node['line'] = head.sourceline


def set_outline_level(p, level):
    """
    设置文档段落的大纲级别
    :param p: 段落
    :param level: 大纲级别，1~9，0为正文文本
    """
    p_element = p._p
    p_pr = p_element.find(qn('w:pPr'))
    if p_pr is None:
        p_pr = oxml.OxmlElement('w:pPr')
        p_element.insert(0, p_pr)
    outline_lvl = p_pr.find(qn('w:outlineLvl'))
    if outline_lvl is not None:
        p_pr.remove(outline_lvl)
    new_outline_lvl = oxml.OxmlElement('w:outlineLvl')
    new_outline_lvl.set(qn('w:val'), str(min(8, level - 1)) if level else '-1')
    p_pr.append(new_outline_lvl)


def output_doc(juan, doc, book, types):
    """
    向文档输出疏钞内容
    :param juan: 疏钞的卷号
    :param doc: 文档对象
    :param book: 疏钞HTML的主内容或DIV的soup对象
    :param types: 临时记录段落类型的集合对象，可含 shu、chao、jing、xu
    :return: 生成的段落个数
    """
    count = 0
    for p in book.children:
        p_types, dp = set(types), None
        p_cls = p.attrs.get('class', []) if hasattr(p, 'attrs') else []
        if 'shu' in p_cls:
            p_types.add('shu')
        if 'chao' in p_cls:
            p_types.add('chao')
        if 'jing' in p_cls:
            p_types.add('jing')
        if 'xu' in p_cls:
            p_types.add('xu')
        if (p.name == 'p' or 'text_paragraph' in p_cls or 'juan-begin' in p_cls) and p.text:
            if 'head' in p_cls and '2.1.3 約所詮辨異' in p.text:
                add_final_vol(doc)
            dp = output_p(doc, p, p_types)
            count += 1
        elif p.name == 'div':
            sub_n = output_doc(juan, doc, p, p_types)
            if sub_n:
                count += sub_n
            elif p.text:
                dp = output_p(doc, p, p_types)
                count += 1

        if dp:
            pick_0236(juan, doc, dp)
            mark_shu_chao_underline(dp)

    return count


def mark_shu_chao_underline(dp):
    """根据疏钞段落的【标记】设置下划线"""
    m = re_p_tag.search(dp.text)
    tag = m.group() if m else ''
    text = dp.text[len(tag):]
    dp.clear()
    if data['test'] and dp.style.name[0] in '疏钞经':
        dp.add_run(f"{data['p_n']:04d}{'' if tag else ' '}", '段号')
    if tag:
        dp.add_run(tag)
    add_runs_u(dp, text, data['test'] and data.get('mark_shu'), False)


def output_p(doc, el, types):
    """
    向文档输出一个疏钞段落
    :param doc: 文档对象
    :param el: 疏钞HTML的段落节点
    :param types: 段落类型的集合对象
    :return: 添加的文档段落对象
    """
    p_cls = el.attrs.get('class', [])
    kp = 'head' in p_cls
    p_style = '科判' if kp else '卷' if 'juan-begin' in p_cls else ''
    if not p_style and types:
        if 'shu' in types:
            p_style = '疏'
        elif 'chao' in types:
            p_style = '钞'
        elif 'jing' in types or re_jing.search(el.text):
            p_style = '经'

    for r in el.find_all('ruby'):
        r.string = f'({r.text.strip()})'
    text = fix_kuo(el.text)

    level = kp and re_kp_level.search(text)
    m = re_p_tag2.search(text)
    if m:
        c = m.group()[0]
        assert {'鈔': '钞', '經': '经'}.get(c, c) == p_style
        text = '【' + m.group(1) + '】' + text[len(m.group()):]
    elif level:
        text = re.sub(r'^\d+:', '', re.sub(r'^\d+:…', lambda _: _.group().replace(':', ' '), text))

    # assert p_style or text.endswith('(终)')
    # assert not re.search('[「《].+[」》]作「', text)

    p = doc.add_paragraph(text, p_style or '段落')
    if level:
        set_outline_level(p, int(level.group()))
    if 'juan-end-end' in p_cls:
        set_outline_level(p, 1)
        p.alignment = doc_t.WD_PARAGRAPH_ALIGNMENT.RIGHT
    el.string = ''
    return p


def fix_kuo(text):
    """修正段落文本中的括号"""
    text = re_kuo_dup.sub(lambda _: '「' + re_kuo_sign.sub('', _.group()) + '」', text.strip())
    text = re_kuo_embed.sub(fix_kuo_embed, text)
    text = re_kuo_long.sub(lambda _: _.group().replace('「', '“').replace('」', '”'), text)
    text = re_kuo_end.sub('”', re_kuo_sign.sub('', text))
    return text


def fix_kuo_embed(m):
    """修正嵌套引号的文本"""
    text = re.sub('『[^『』]{0,4}』', lambda _: _.group()[1: -1], m.group())
    return text.replace('『', '‘').replace('』', '’').replace('「', '“').replace('」', '”')


def add_final_vol(doc):
    """添加下册分隔段落"""
    p = doc.add_paragraph('华严经疏钞附会玄记合编上册(终)', '疏')
    p.alignment = doc_t.WD_PARAGRAPH_ALIGNMENT.RIGHT

    p = doc.add_paragraph('华严经疏钞附会玄记合编 下册', '科判')
    p.alignment = doc_t.WD_PARAGRAPH_ALIGNMENT.CENTER
    set_outline_level(p, 1)

    p = doc.add_paragraph('7 …2.2.3.2.3 立教開宗', '科判')
    set_outline_level(p, 2)


def pick_0236(juan, doc, p):
    """
    根据刚添加的疏钞段落对象，添加会玄记段落
    :param juan: 疏钞的卷号
    :param doc: 文档对象
    :param p: 疏钞段落对象
    :return: 生成的会玄记段落个数
    """
    scj = p and p.style.name[0] in '疏钞经'
    if p:
        if scj and data.get('last_p_type') != p.style.name and not p.text.startswith('【'):
            p.text = f'【{p.style.name}】' + p.text  # 类型改变后加段首标注
        if scj and '会' in data.get('last_p_type', ''):  # 疏钞与前面的会隔开
            p.insert_paragraph_before('', '空行')
        data['last_p_type'] = data['last_type0'] = p.style.name
    if not scj:
        return 0
    data['p_n'] += 1  # p_n、last_type0 为当前疏钞经的段号、类型

    p_text, text0 = p.text, re_punc.sub('', p.text)
    mode, mis_i, count, first_i = 0, 0, 0, 0

    # 从第一个未合入的会段开始，尝试在当前疏钞经下合入
    count, _ = pick_loop(doc, p, juan, 4, mis_i, first_i, count, text0)
    count, first_i = pick_loop(doc, p, juan, mode, mis_i, first_i, count, text0)
    if count and first_i and 0:  # 在csv中插入疏钞段落
        p = dict(type0=p.style.name[0], type='', id=0, head='', juan='',
                 juan0=juan, para=data['p_n'], text=re_p_tag.sub('', p_text))
        data['pu'].insert(first_i - 1, p)

    return count


def pick_loop(doc, p, juan, mode, mis_i, first_i, count, text0):
    """pick_0236的内部循环"""
    orig_num = 0
    for i, pu in enumerate(data['pu']):
        if mode == 0 and pu['para'] == 0:
            mode = 1
        if mode == 2 and pu['para'] == 0:
            if i > mis_i + 20:
                mode = 3
            elif pu.get('force'):  # 遇到指定了匹配文本的段落
                org = match_ref(juan, doc, p, pu, text0, 1)[-1]  # 仅检测匹配
                if org:  # 继续合入后续段落
                    mode = 1
        if mode == 1 and pu['para'] == 0:
            dp, ts, org = match_ref(juan, doc, p, pu, text0)  # 尝试合入
            if dp:
                count += 1
                first_i = first_i or i + 1
            elif not org:  # 不匹配则忽略后续段，直到遇到指定了匹配文本的段落再合入后续段落
                mode = 2
                mis_i = i
        elif mode == 4 and pu['para'] == 0 and pu.get('orig') and (
                not orig_num or orig_num == pu['orig']):
            if orig_num or pu.get('force'):
                dp, ts, org = match_ref(juan, doc, p, pu, text0)
                if dp:
                    count += 1
                    orig_num = pu['orig']

    return count, first_i


def match_ref(juan, doc, p, pu, text0, test=0):
    """
    匹配当前疏钞内容，创建会玄记段落
    :param juan: 疏钞的卷号
    :param doc: 文档对象
    :param p: 疏钞段落对象
    :param pu: 会玄记段落数据
    :param text0: 疏钞段落文本，无标点
    :param test: 创建模式，0-匹配并创建，1-匹配但不创建，2-直接创建
    :return: (dp, ts, word0) 创建的文档段落对象、各个可能的关键词、第一个匹配的关键词
    """
    first_w, orgs, exists, ts = find_terms(p, pu, text0)

    dp = add_pu_para(juan, doc, p, pu, orgs, exists, test, first_w)
    [orgs.remove(s) for s in exists if s in orgs]

    return dp, ts, orgs and orgs[0]


def find_terms(p, pu, text0):
    """查找匹配词组"""
    orgs, exists, txt0, pu_t = [], [], text0, pu['text']
    force = pu.get('force', [])
    a = re_tag1.findall(pu_t)
    b = re_tag2.findall(pu_t)
    c = re_tag3.findall(pu_t)
    ts = [str(t).strip(punc_tail)
          for t in sorted(a + b + c, key=len, reverse=True) if len(t) < 20]
    min_i, first_w, tps = 0, '', []

    for t in ts:
        tp = t
        if force and t in force:
            si = force.index(t)
            if si % 2 == 0:
                tp = force[si + 1]
        t1 = re_term.sub(lambda _: _.group()[1:-1], t)
        if t1 != t and t1 in ts and tp == t:
            continue
        s = len(t) > 1 and t not in orgs and t not in pu.get('skip', []) and re_punc.sub('', tp)
        if s and s in text0:
            tps.append((t, tp, s))
            orgs.append(t)
            i = pu_t.index(t)
            if min_i > i or not first_w:
                min_i = i
                first_w = t
    ts = [t for t in tps if t[0] == first_w] + [t for t in tps if t[0] != first_w]
    orgs = []
    for t, tp, s in ts:
        exist = s not in txt0 or '【' + tp in p.text or tp + '】' in p.text
        if not exist and s in txt0:  # 在疏钞中也有则加到orgs，add_pu_para 中再匹配force中的疏钞文本
            orgs.append(t)
            txt0 = txt0.replace(s, '')
        elif exist:
            orgs.append(t)
            exists.append(t)
    return first_w, orgs, exists, ts


def add_pu_para(juan, doc, p, pu, orgs, exists, test=0, first_w=''):
    """
    创建会玄记段落
    :param juan: 疏钞的卷号
    :param doc: 文档对象
    :param p: 疏钞段落对象
    :param pu: 会玄记段落数据
    :param orgs: 所有已匹配的关键词
    :param exists: 疏钞段落中不能再接受的关键词，同时存在于orgs
    :param test: 创建模式，0-匹配并创建，1-匹配但不创建，2-直接创建
    :param first_w: 会玄记段落中第一个匹配的关键词
    :return: 创建的文档段落对象
    """
    force = pu.get('force', [])
    if force:
        update_terms_by_force(force, orgs, p, pu)
    if pu.get('skip'):
        [orgs.remove(s) for s in pu['skip'] if s in orgs]

    if (orgs and pu['para'] != data['p_n'] and not test) or test == 2:
        style = pu['type']
        assert style == {'鈔': '钞', '經': '经'}.get(style, style)
        if pu['para'] == 0 and set(orgs) - set(exists):
            pu_text, pre = pu['text'], ''
            first_miss = first_w and (first_w not in orgs and first_w in exists)
            if first_miss and first_w not in orgs:
                orgs.insert(0, first_w)
            for org in list(orgs):
                if not force and ('【' + org in p.text or org + '】' in p.text):
                    exists.append(org)
                if org not in exists:
                    p.text = replace_term(p.text, org)
                s = f'言{org}者'
                s = s if s in pu_text else org
                if org in exists and s not in exists:
                    exists.append(s)
                pu_text = replace_term(pu_text, s)

            if not orgs and test != 2:
                return
            return add_pu_p(doc, p, pu, orgs, exists, juan, pu_text, pre, style, first_miss, force, test)


def add_pu_p(doc, p, pu, orgs, exists, juan, pu_text, pre, style, first_miss, force, test):
    """将一个会玄记段落添加到文档"""
    if force and not test:
        for i, s in enumerate(force):
            if i % 2 == 1:
                p.text = replace_term(p.text, s)
    if data['last_p_type'] != style + '会':
        pre = f"【会{'' if data['last_p_type'].replace('会', '') == style else style}】"
        if force and data['last_p_type'][0] != style:
            print(f"{style} != {data['last_p_type']}", pu['head'], style, ','.join(force), data['p_n'])
    data.update(dict(last_p_type=style + '会', last_pu_id=pu['id']))
    pu.update(dict(para=data['p_n'], juan0=juan,
                   words=orgs, type0=data['last_type0']))
    dp = doc.add_paragraph('', '会')
    if data['test']:
        head = '_' + pu['head'] and ''
        dp.add_run(f"{pu['id']:04d}{head}{'' if pre else ' '}", '段号-会')
    dp.add_run(pre)
    pu_text = re.sub('▲(。|[疏鈔]：)+', '▲', pu_text)
    add_runs_u(dp, re.sub('^▲', '', pu_text) if pre else pu_text,
               bold=data['test'], first_miss=first_miss, exists=exists)
    if first_miss:
        pu['first_miss'] = orgs.pop(0)
    if pu.get('has_kp'):
        dp.add_run('(科判略)', 'note')
    return dp


def update_terms_by_force(force, orgs, p, pu):
    """ 根据会玄记中p元素的force属性(匹配文本)更新当前匹配词组 """
    if len(force) % 2 != 0:
        assert 0, f"{pu['head']} force 文本两个一组，逗号分隔"
    ins_i = 0
    for i, s in enumerate(force):
        if i % 2 == 0:
            assert s and force[i + 1], f"{pu['head']} force 两个文本都不能空"
            assert s in pu['text'], f"{pu['head']} force 第一个文本要在X0236段中存在"
            if force[i + 1] in p.text:
                s = s.strip(punc_tail)
                if s not in orgs:
                    orgs.insert(ins_i, s)
                    ins_i += 1
            else:
                s = re_punc.sub('', force[i + 1])
                if s in re_punc.sub('', p.text) and not re.search('^[“○。]|[；。，、]$', force[i + 1]):
                    print(f"{pu['head']} 的 force 标点可能需修正", data['p_n'], force[i + 1] + '\n',
                          p.text[p.text.index(s[0]):][:150])
                orgs.clear()
                break


def replace_term(text, w):
    """
    在段落文本中用【】标记已匹配的关键词
    :param text: 段落文本
    :param w: 关键词
    :return: 段落文本
    """
    m = re.search(f'(^|[^【]){w}($|[^】])', text)
    return text.replace(m.group(), m.group().replace(w, f'【{w}】'), 1) if m else text


def add_runs_u(p, text, underline=True, bold=True, first_miss=False, exists=None):
    """
    根据段落文本中的【】标记，向文档段落对象添加普通文本和下划线文本
    :param p: 文档段落对象
    :param text: 段落文本
    :param underline: 是否加下划线
    :param bold: 是否对下划线文本加粗
    :param first_miss: 会玄记段落中第一个关键词是否与疏钞段落匹配
    :param exists: 疏钞段落中不能再接受的关键词
    """
    texts = re_p_tags.split(text)
    n, last = 0, None
    for i, txt in enumerate(texts):
        if i:
            j = text.index('】', n)
            s = text[n + 1: j]
            if re_term.match(s):
                p.add_run('言')
                r = p.add_run(s[1:-1])
                last = p.add_run('者')
            else:
                if last and last.underline:
                    p.add_run(' ').font.size = Pt(7)
                r = last = p.add_run(s)
            if underline:
                if first_miss:
                    first_miss = False
                    r.italic = True
                r.underline = True
                r.bold = bold and not (exists and s in exists)
            n = j + 1
        if txt:
            cjk_b, cjk_c = re_cjk_b.findall(txt), re_cjk_c.findall(txt)
            if cjk_b or cjk_c:
                s = ''
                for c in list(txt):
                    if c in cjk_b or c in cjk_c:
                        if s:
                            last, s = p.add_run(s), ''
                        if c in cjk_c:
                            p.add_run(c, 'cjk-c')
                        elif c in cjk_b:
                            p.add_run(c, 'cjk-b')
                    else:
                        s += c
                if s:
                    last = p.add_run(s)
            else:
                last = p.add_run(txt)
        n += len(txt)


def load_0236(html_dir):
    """
    加载各卷会玄记的段落内容
    :param html_dir: 会玄记网页的目录
    :return: 会玄记段落节点数据的数组
    """
    pu = []
    re_juan = re.compile(r'_(\d+)\.html')
    data['orig_n'] = 0
    for fn in sorted(glob(html_dir + '/*.html')):
        juan = re_juan.search(fn)
        juan = int(juan.group(1)) if juan else -1
        if juan < 0 or (data['test'] and juan < data.get('start_juan', 0)):
            continue
        with open(fn, encoding='utf-8') as f:
            soup = BeautifulSoup(f.read().replace('[＊]', '').replace(
                '》、《', '》《'), 'html.parser')
        body = soup.find('div', id='body')
        for s in (body.select('.no-marker') or []):
            p = s.previous_sibling
            while p and p.name != 'p':
                p = p.previous_sibling
            if p and p.name == 'p':
                p['has_kp'] = True
        for s in body.select('p.ignore'):
            head = s.find('span', class_='lineInfo').attrs['line']
            s.string = '-' + head
        for s in body.select('.juan,.byline,.no-marker,.div-w'):
            s.extract()
        start = soup.find('span', dict(line='0094b07')).sourceline if juan == 1 else 1
        load_0236_div(body, start, dict(juan=juan), pu)

    heads = sorted(list(set([p['id'] for p in pu])))
    start_id = data.get('start_id', 1) if data['test'] else 1
    for p in list(pu):
        p['id'] = start_id + heads.index(p['id'])
        if len(p['text']) < 2 or p['text'][0] == '-' or p['para'] < 0:
            pu.remove(p)
    return pu


def load_0236_div(body, start, status, pu):
    """
    解析会玄记网页的一个DIV
    :param body: DIV节点
    :param start: 起始行号，段落行号大于此值的才解析
    :param status: 上下文对象
    :param pu: 存储会玄记段落节点数据的数组
    """
    for p in body.children:
        p_cls = p.name and p.attrs.get('class', [])
        if p.name == 'div':
            status1 = dict(status)
            if 'div-orig' in p_cls:
                data['orig_n'] += 1
                status1.update(dict(p_type='疏', orig=data['orig_n']))
            elif 'div-commentary' in p_cls:
                status1['p_type'] = '钞'
            load_0236_div(p, start, status1, pu)
        elif p.name == 'p' and p.text and p.sourceline > start:
            load_0236_p(status, p_cls, pu, p)


def load_0236_p(status, p_cls, pu, p):
    """ 解析会玄记网页的一个段落 """
    text = re.sub(r'\[[A-Z0-9]+]', '', p.text.strip())
    m = re.search('^[▲○]?([疏鈔钞])云?[：「]', text)
    status['p_type'] = m.group(1) if m else status.get('p_type', '钞')
    if 'shu' in p_cls:
        status['p_type'] = '疏'
    elif 'chao' in p_cls or status['p_type'] == '鈔':
        status['p_type'] = '钞'
    if status['p_type'] == '疏':
        text = re.sub('^[▲○]?疏[云：▲○]+', '', text)
    elif status['p_type'] == '钞':
        text = re.sub('^[▲○]?[鈔钞][云：▲○]+', '', text)
    else:
        text = re.sub('^[▲○]?[疏鈔钞][云：▲○]+', '', text)
    text = text.replace('「', '“').replace('」', '”').replace('『', '‘').replace('』', '’')
    text = re.sub('\n▲[疏鈔]：', '\n▲', text)

    m = re.search('^-([a-d0-9]+)', text)
    if m:
        head = m.group(1)
    else:
        head = p.find('span', class_='lineInfo').attrs['line']

    m = re.search(r'^(\d{4})([abcd])(\d{2})$', head)
    hid = int(m.group(1)) * 1000 + 'abcd'.index(m.group(2)) * 100 + int(m.group(3))
    item = dict(type=status['p_type'], juan=status['juan'],
                head=head, id=hid, para=0, text=text)
    pu.append(item)
    if 'skip' in p.attrs:
        item['skip'] = p.attrs['skip'].split(',')
    if 'force' in p.attrs:
        item['force'] = p.attrs['force'].split(',')
    if 'has_kp' in p.attrs:
        item['has_kp'] = True
    if status.get('orig'):
        item['orig'] = status['orig']
    if data.get('diff_hui'):
        diff_hui_p(pu, status['juan'], head, text)

    for span in p.find_all('span', class_='lineInfo'):
        head_ = span.attrs['line']
        if head_ != head:
            next_s = span.next_sibling
            assert next_s and next_s.text.startswith('▲')
            m = re.search(r'^(\d{4})([abcd])(\d{2})$', head_)
            hid = int(m.group(1)) * 1000 + 'abcd'.index(m.group(2)) * 100 + int(m.group(3))
            pu.append(dict(type=status['p_type'], juan=status['juan'],
                           head=head_, id=hid, para=-1, text='-'))


def diff_hui_p(pu, juan, head, text):
    """将所用的会玄记段落与CB最新下载的网页作对比"""
    found, text_may, re_punc_ = 0, '', re.compile('[“”]')
    for i, r in enumerate(data['origin']):
        r = r or [0, 0, 0, '']
        r[-1] = r[-1].replace('“慨”者', '“慨然”者').replace('言關中', '言關脈中')
        if r[1] == '0192a03' and r[2] == '15':
            r[2] = '14'
        text1 = re_punc_.sub('', r[-1])
        if r[1] == head and int(r[2]) == juan:
            found = found or 1
            text_may = r[-1]
            if found == 1 and text1[:10] == re_punc_.sub('', text)[:10]:
                found = 2
                break
    if found != 2:
        print(diff_text(text_may, text), text_may)
        print(len(pu), pu[-1])
        exit(1)


def save_used(pu):
    """
    保存会玄记匹配结果到 pu-used.csv
    :param pu: 会玄记段落节点数据的数组
    """
    with open('pu-used.csv', 'w', encoding='utf-8') as f:
        f.write('段号,行首,会卷,移动,会标,疏卷,疏段,疏标,匹配数,警告,首词不匹配,字数,会玄记文本\n')
        re_trim = re.compile(r'^[▲○]?[疏鈔钞]云?[：「]|\n.+$', re.M)
        para0, h0 = 0, 0
        for r in pu:
            h, para, words = r['id'], r['para'], r.get('words', [])
            t0, t = r.get('type0', ''), r['type']
            move = t and f'上移{h0 - h}' if h < h0 else f'下移{h - h0 - 1}' if h > h0 + 1 else ''
            text = re_trim.sub('', r['text']) + ('(科判略)' if r.get('has_kp') else '')
            warn = '逆序' if para0 - 5 > para > 0 else '跳变' if 0 < para0 < para - 10 else ''
            f.write(f"{h or ''},{r['head']},{r['juan']},{move},{t},"
                    f"{r.get('juan0', '未合')},{para or '未合'},"
                    f"{t0 and t0 + ('' if t0 == t or not t else '异')},"
                    f"{t and (len(words) or '无匹配')},{t and warn},"
                    f"{'□' + r['first_miss'] if r.get('first_miss') else ''},"
                    f"{len(re_punc.sub('', text))},{text[:80]}\n")
            para0 = para or para0
            h0 = h


def main(release=0, min_juan0=0, max_juan0=10, kp_build=0, gen_origin=0, diff_hui=0, mark_shu=1):
    """
    文档生成的主函数
    :param release: 是否生成发布用的文档
    :param min_juan0: 是否生成校对文档的起始卷号
    :param max_juan0: 是否生成校对文档的结束卷号
    :param kp_build: 是否重新生成科判json文件
    :param gen_origin: 是否根据下载的 X0236 网页生成会玄记原始内容文件 pu-origin.csv
    :param diff_hui: 是否根据会玄记原始内容文件 pu-origin.csv 对比当前 hui 下的内容
    :param mark_shu: 生成校对文档时是否将会玄记引用词在疏钞段中匹配文本加下划线
    """
    data.update(dict(test=not release, gen_origin=gen_origin, diff_hui=diff_hui,
                     min_juan0=min_juan0, max_juan0=max_juan0, mark_shu=mark_shu))
    if data.get('diff_hui'):
        with open('pu-origin.csv', encoding='utf-8') as f:
            data['origin'] = [r.strip().split(',') for r in f.readlines()[1:]]
        if load_0236('./hui'):
            return
    kp_map = {'tree': [], 'id2node': {}, 'heads': set()}
    doc = Document('hysc.docx')

    if not kp_build:
        with open('hysc.json', encoding='utf-8') as f:
            kp_map = json.load(f)
        data['p_n'] = data.get('shu_para0', 1) - 1 if data['test'] else 0
        data['pu'] = load_0236('./hui')

    for fn in sorted(glob('./shu/*.html')):
        juan = int(re.search(r'juan(\d+)\.html', fn).group(1))
        if data['test'] and juan and juan not in range(min_juan0, max_juan0):
            continue
        with open(fn, encoding='utf-8') as f:
            soup = re_kuo_kuo.sub(lambda m: m.group(2) + m.group(3), f.read())
            soup = BeautifulSoup(soup.replace('》、《', '》《'), 'html.parser')
        if juan == 0 and kp_build:
            try:
                read_kp(soup, kp_map)
            except AssertionError:
                pass
        book = soup.find('div', id='main')
        add_head_level(kp_map, book, str(juan))
        if juan > 0:
            for s in book.select('.juan-begin>p,.juan-begin>.text_paragraph,.juan-end,.noteAnchor,.double-circle'):
                s.extract()
            output_doc(juan, doc, book, set())

    if kp_build:
        with open('hysc.json', 'w', encoding='utf-8') as f:
            json.dump(dict(tree=kp_map['tree'], id2node=kp_map['id2node']), f, ensure_ascii=False, indent=2)
    else:
        doc.save(f"合编{'' if data['test'] else '-'}.docx")
        save_used(data['pu'])


if __name__ == '__main__':
    fire.Fire(main)
